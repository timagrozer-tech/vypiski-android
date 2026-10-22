"""12.79.9: a § of unnumbered blocks (one per serviceman, each closed by «Основание:») — the excerpt holds only the person's block;
DOCX carries «do not expand lines ending with a manual break»."""
import sys as _s,os;_s.path.insert(0,os.path.dirname(os.path.abspath(__file__)));from authhelp import login
import asyncio, sys, http.server, threading, functools, io, zipfile
from playwright.async_api import async_playwright
from docx import Document
ROOT=sys.argv[1];OUT='/tmp/vypiski-test-out/blocks/';os.makedirs(OUT,exist_ok=True)
d=Document()
for t in ['ПРИКАЗ','КОМАНДИРА ВОЙСКОВОЙ ЧАСТИ 11111','по строевой части','« 05 » декабря 2025 г. № 361','с. Энское','§ 1']:d.add_paragraph(t)
for i,(fam,fio) in enumerate([('ПЕРВОВОЙ Анне Ивановне – матери','ПЕРВОВА Ивана Петровича, АА-111111'),('ВТОРОВОЙ Ольге Петровне – матери','ВТОРОВА Петра Ивановича, ББ-222222'),('ТРЕТЬЕВОЙ Нине Павловне – жене','ТРЕТЬЕВА Олега Сергеевича, ВВ-333333')]):
  d.add_paragraph(f'В соответствии с Указом Президента Российской Федерации от 25 июля 2006 года № 765 выплатить единовременное пособие в размере пяти окладов денежного содержания {fam} рядового {fio}, стрелка, исключенного из списков личного состава в связи со смертью:')
  d.add_paragraph('единовременное поощрение в размере пяти окладов денежного содержания при награждении ОРДЕНОМ МУЖЕСТВА посмертно')
  d.add_paragraph(f'Мать: получатель, паспорт серия 0000 № 00000{i}.')
  d.add_paragraph(f'Основание: Указ Президента Российской Федерации от 21.05.2025 года №33{i}сс.')
for t in ['','ВрИО КОМАНДИРА ВОЙСКОВОЙ ЧАСТИ 11111','гвардии подполковник','С.Командиров']:d.add_paragraph(t)
d.save(OUT+'Приказ 361.docx')
h=functools.partial(http.server.SimpleHTTPRequestHandler, directory=ROOT); h.log_message=lambda *a:None
srv=http.server.ThreadingHTTPServer(("127.0.0.1",0),h); threading.Thread(target=srv.serve_forever,daemon=True).start()
res=[]
def check(n,ok,dd=''):res.append(bool(ok));print(('PASS ' if ok else 'FAIL ')+n+(' — '+str(dd)[:300] if dd else ''))
async def main():
  async with async_playwright() as p:
    b=await p.chromium.launch();pg=await (await b.new_context()).new_page();errs=[];pg.on("pageerror",lambda e:errs.append(str(e)))
    await pg.goto(f"http://127.0.0.1:{srv.server_address[1]}/index.html");await login(pg)
    await pg.evaluate("()=>VypiskiShell.go('import')");await pg.wait_for_timeout(500)
    await pg.set_input_files("#wordFiles",[OUT+'Приказ 361.docx']);await pg.wait_for_timeout(1500)
    await pg.evaluate("()=>[...document.querySelectorAll('#content button.primary')].find(x=>/Импорт/.test(x.textContent)).click()")
    await pg.wait_for_timeout(500)
    while await pg.evaluate("()=>VypiskiImportJob.state().running"):await pg.wait_for_timeout(300)
    r=await pg.evaluate("""async()=>{const m=(await VypiskiCore.allMeta())[0];const a=await VypiskiWizard.auto({orderId:m.orderId,query:'Второв Петр Иванович',full:'Второв Петр Иванович',mode:'fio',refs:[{s:'1',p:''}]});
      const rows=await VypiskiRegisterView.scan(r=>/ВТОРОВ П/.test(r.fio));const blob=await VypiskiDocument.toDocx(a.model,a.t);const z=await JSZip.loadAsync(await blob.arrayBuffer());
      return {ok:a.ok,fr:a.fragments.map(f=>f.text),cat:rows.map(r=>r.category),settings:await z.file('word/settings.xml')?.async('string')||''}}""")
    check('В выписке только блок этого человека',r['ok'] and len(r['fr'])==5 and 'ВТОРОВОЙ' in r['fr'][1] and not any('ПЕРВОВ' in x or 'ТРЕТЬЕВ' in x for x in r['fr']),r['fr'])
    check('Вид события — выплата, а не гибель',r['cat']==['Выплаты и надбавки'],r['cat'])
    check('DOCX: строки с ручным переносом не растягиваются','doNotExpandShiftReturn' in r['settings'],r['settings'][:120])
    check('Без ошибок',not errs,errs);await b.close()
  print(f"{sum(res)}/{len(res)}")
asyncio.run(main())
