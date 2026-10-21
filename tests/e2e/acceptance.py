import os as _o, sys as _s;_s.path.insert(0,_o.path.dirname(_o.path.abspath(__file__)));from authhelp import login
"""Acceptance checks from DESIGN_SPEC 2.0 (no builds): run against the page in a browser and in a simulated Android app."""
import os
FIX=os.environ.get("VYP_FIX","fixtures-out")
OUTDIR=os.environ.get("VYP_OUT","/tmp/vypiski-test-out");os.makedirs(OUTDIR,exist_ok=True)

import asyncio, sys, os, re, json, glob, hashlib, http.server, threading, functools, zipfile
from playwright.async_api import async_playwright
ROOT=sys.argv[1]; OUT=f"{OUTDIR}/accept"; os.makedirs(OUT,exist_ok=True); ZIP=f"{FIX}/accept/приказы.zip"
sha=lambda p: hashlib.sha256(open(p,'rb').read()).hexdigest()
JOBS={}
class H(http.server.SimpleHTTPRequestHandler):
    def __init__(self,*a,**k): super().__init__(*a,directory=ROOT,**k)
    def log_message(self,*a): pass
    def do_GET(self):
        m=re.match(r"^/saf/([^/]+)/(\d+)$",self.path)
        if not m: return super().do_GET()
        data=open(JOBS[m.group(1)][int(m.group(2))],'rb').read(); rng=self.headers.get("Range")
        if rng:
            a,b=re.match(r"bytes=(\d+)-(\d*)",rng).groups();a=int(a);b=int(b) if b else len(data)-1
            self.send_response(206);self.send_header("Content-Range",f"bytes {a}-{b}/{len(data)}");body=data[a:b+1]
        else: self.send_response(200);body=data
        self.send_header("Content-Length",str(len(body)));self.end_headers();self.wfile.write(body)
srv=http.server.ThreadingHTTPServer(("127.0.0.1",0),H);threading.Thread(target=srv.serve_forever,daemon=True).start()
URL=f"http://127.0.0.1:{srv.server_address[1]}/index.html"
results=[]
def check(name,ok,detail=''):
    results.append((name,bool(ok)));print(('PASS ' if ok else 'FAIL ')+name+(' — '+str(detail) if detail else ''))

async def wait_import(pg):
    for _ in range(400):
        await pg.wait_for_timeout(250)
        if not await pg.evaluate("()=>VypiskiImportJob.state().running"): break
    return await pg.evaluate("()=>{const j=VypiskiImportJob.state();return {st:j.result&&j.result.st,errors:j.errors.map(e=>e.path+': '+e.error)}}")

async def main():
    zip_sha_before=sha(ZIP); src101=f'{FIX}/orders/Приказ_101_2016.docx'; src_sha_before=sha(src101)
    async with async_playwright() as p:
        b=await p.chromium.launch()
        # ---------- 1. Browser / Windows: ZIP with «приказы/год/месяц», hostile entries ----------
        ctx=await b.new_context(viewport={"width":1366,"height":900},locale="ru-RU",accept_downloads=True); pg=await ctx.new_page()
        errs=[];pg.on("pageerror",lambda e:errs.append(str(e)));pg.on("dialog",lambda d:asyncio.ensure_future(d.accept()))
        await pg.goto(URL);await pg.wait_for_timeout(1500)
        await login(pg)
        await pg.evaluate("()=>document.querySelector(`[data-sec='import']`).click()");await pg.wait_for_timeout(600)
        await pg.set_input_files("#wordFiles",ZIP);await pg.wait_for_timeout(2500)
        status=await pg.inner_text("#impPrepStatus");struct=await pg.evaluate("()=>{const d=document.querySelector('#content details');return d?d.innerText:''}")
        check("ZIP: найдено 9 документов Word (8 + дубликат), служебный «~$» пропущен","Найдено документов: 9" in status and "«~$» пропущено: 1" in status,status)
        check("ZIP: распознана структура годов 2025, 2026","2025" in struct and "2026" in struct,struct.split('\n')[0])
        await pg.evaluate("()=>document.querySelector('#content details').open=true")
        struct=await pg.evaluate("()=>document.querySelector('#content details').innerText")
        check("ZIP: месяцы январь/февраль/март и корневая папка «приказы»","январь" in struct and "февраль" in struct and "март" in struct and "приказы" in struct)
        warn=await pg.inner_text("#content .imp-log >> nth=-1")
        check("ZIP: Zip Slip-запись отклонена","Небезопасный путь" in warn,[l for l in warn.split('\n') if 'evil' in l])
        check("ZIP: вложенный архив отмечен","вложенных архивов" in warn)
        await pg.evaluate("()=>[...document.querySelectorAll('#content button.primary')].find(x=>/Импорт/.test(x.textContent)).click()")
        r=await wait_import(pg);st=r['st']
        check("ZIP: сохранено 8 приказов, дубликат пропущен, ошибок импорта нет",st['saved']==8 and st['dupFile']+st['dupText']==1 and st['failed']==0,{k:st[k] for k in ('saved','dupFile','dupText','failed')})
        check("ZIP: «бомба» (60 МБ) отклонена ещё при анализе",any('бомба' in e and ('сжат' in e or '60 МБ' in e) for e in r['errors']) and 'бомба' in warn,[e for e in r['errors'] if 'бомба' in e])
        check("ZIP: PDF не импортирован и не выдан за приказ",st['saved']==8 and not await pg.evaluate("async()=>(await VypiskiRepository.all('orders')).some(o=>/\\.pdf$/i.test(o.sourceFile))"))
        check("ZIP: архив на диске не изменён (SHA-256)",sha(ZIP)==zip_sha_before)
        # ---------- 2. SHA-256 before/after + excerpt contents ----------
        order=await pg.evaluate("async()=>{const o=(await VypiskiRepository.all('orders')).find(o=>o.sourceFile==='Приказ_101_2016.docx');return {id:o.id,fileSha256:o.fileSha256,textSha256:o.textSha256,stored:o.originalStored,paras:o.paragraphs.map(p=>p.text)}}")
        check("Оригинал в базе побайтно равен исходному файлу (SHA-256)",order['fileSha256']==src_sha_before,order['fileSha256'][:16])
        person=await pg.evaluate("id=>VypiskiRegister.extract({paragraphs:[]})&&VypiskiRepository.get('orders',id).then(o=>VypiskiRegister.extract(o)[0])",order['id'])
        await pg.evaluate("id=>VypiskiWizard.open({orderId:id,query:'',mode:'fio'})",order['id']);await pg.wait_for_timeout(1200)
        await pg.locator("#wz input >> visible=true").first.fill(person['full'] or person['fio'])
        await pg.locator("#wz button:has-text('Найти в приказе') >> visible=true").first.click();await pg.wait_for_timeout(1000)
        for i in range(8):
            nxt=pg.locator("#wz .wz-next >> visible=true")
            if not await nxt.count(): break
            if await nxt.first.is_disabled():
                h=pg.locator("#wz .wz-why >> visible=true")
                if await h.count():
                    await h.first.click();await pg.wait_for_timeout(300);await pg.evaluate("()=>{const a=document.activeElement;if(a&&a.type==='checkbox'&&!a.checked)a.click()}");await pg.wait_for_timeout(300)
            await nxt.first.click();await pg.wait_for_timeout(900)
        async with pg.expect_download() as d:
            await pg.locator("#wz button:has-text('Скачать выписку DOCX')").first.click()
        dl=await d.value;docx=f"{OUT}/{dl.suggested_filename}";await dl.save_as(docx);await pg.wait_for_timeout(1500)
        final=await pg.inner_text("#wz")
        line=[l for l in final.split('\n') if 'до и после' in l]
        check("Выписка: показано «SHA-256 приказа до и после … совпадает»",line and 'совпадает, оригинал не изменён' in line[0],line[0] if line else final[-300:])
        aud=await pg.evaluate("async()=>(await VypiskiAudit.list({limit:20,type:'excerpt.exported'}))[0].details")
        check("Аудит: SHA-256 файла и текста до == после",aud['originalSha256Before']==aud['originalSha256After']==order['fileSha256'] and aud['textSha256Before']==aud['textSha256After']==order['textSha256'])
        after=await pg.evaluate("id=>VypiskiRepository.get('orders',id).then(o=>[o.fileSha256,o.textSha256,o.paragraphs.map(p=>p.text)])",order['id'])
        check("Приказ в базе не изменился после генерации",after[0]==order['fileSha256'] and after[1]==order['textSha256'] and after[2]==order['paras'])
        check("Исходный файл на диске не изменился (SHA-256 до/после)",sha(src101)==src_sha_before)
        check("Выписка — отдельный файл .docx",docx.endswith('.docx') and os.path.getsize(docx)>2000 and sha(docx)!=src_sha_before,os.path.basename(docx))
        import docx as pydocx
        text='\n'.join(p.text for p in pydocx.Document(docx).paragraphs)
        sel=[l for l in order['paras'] if (person['fio'].split()[0][:5].lower() in l.lower())]
        other=[l for l in order['paras'] if re.match(r'^\d+\.',l) and l not in sel and person['fio'].split()[0][:5].lower() not in l.lower()]
        check("Выписка содержит выбранный пункт дословно",all(l.strip() in text for l in sel if re.match(r'^\d+\.',l)),sel[:1])
        check("Выписка не содержит невыбранных пунктов",not any(l.strip() in text for l in other),other[:1])
        check("Выписка: заголовок «ВЫПИСКА ИЗ ПРИКАЗА» и реквизиты","ВЫПИСКА ИЗ ПРИКАЗА" in text and "101" in text)
        # 7 screens open without errors
        for sec in ['home','orders','search','import','create','ready','table']:
            await pg.evaluate("s=>document.querySelector(`[data-sec='${s}']`).click()",sec);await pg.wait_for_timeout(700)
        await pg.evaluate("()=>document.querySelector(`[data-sec='orders']`).click()");await pg.wait_for_timeout(1200)
        check("Реестр: панель просмотра справа на широком экране",await pg.is_visible("#u2Pane .u2-integrity"))
        await pg.evaluate("()=>document.querySelector(`[data-sec='ready']`).click()");await pg.wait_for_timeout(800)
        await pg.click("#content button:has-text('История')");await pg.wait_for_timeout(600)
        check("Готовые выписки: история версий",'Создана и сохранена' in await pg.inner_text("#content"))
        check("Нет ошибок JavaScript (браузер / Windows)",not errs,errs[:3])
        await ctx.close()
        # ---------- 3. Android: ZIP chosen through the system picker (SAF), read in ranges ----------
        JOBS['Z1']=[ZIP]
        stub="""(()=>{let pend=null;window.VypiskiNative={version:()=>'12.75',keepAwake:()=>{},exportFolder:()=>'Загрузки / Выписки',resetExportFolder:()=>'',textZoom:()=>100,setTextZoom:()=>{},setDarkBars:()=>{},
         getEntries:(job,from,n)=>JSON.stringify(window.__E.slice(from,from+n)),releaseJob:()=>{},takeShared:()=>{const r=pend;pend=null;return r||''},
         importBegin:()=>{},importProgress:()=>{},importEnd:()=>{},saveBegin:()=>'s',saveChunk:()=>true,saveEnd:()=>'{"ok":true,"where":"x"}',printHtml:()=>{},pickFolder:()=>{},chooseExportFolder:()=>{}};
         window.__share=()=>{pend=JSON.stringify({job:'Z1',count:1,files:1,dirs:0,locks:0,other:0,root:'приказы.zip'});VypiskiNativeEvents.shared()}})();"""
        ctx=await b.new_context(viewport={"width":390,"height":844},locale="ru-RU");await ctx.add_init_script(f"window.__E={json.dumps([['приказы.zip',os.path.getsize(ZIP),0]])};"+stub)
        pg=await ctx.new_page();errs=[];pg.on("pageerror",lambda e:errs.append(str(e)));pg.on("dialog",lambda d:asyncio.ensure_future(d.accept()))
        await pg.goto(URL);await pg.wait_for_timeout(1500)
        await login(pg)
        await pg.evaluate("()=>window.__share()");await pg.wait_for_timeout(3000)
        status=await pg.inner_text("#impPrepStatus")
        check("Android (SAF): ZIP прочитан по частям, найдено 9 документов","Найдено документов: 9" in status,status)
        await pg.evaluate("()=>[...document.querySelectorAll('#content button.primary')].find(x=>/Импорт/.test(x.textContent)).click()")
        r=await wait_import(pg);st=r['st']
        check("Android (SAF): сохранено 8, дубликат и «бомба» отсеяны",st['saved']==8 and st['dupFile']+st['dupText']==1 and any('бомба' in e for e in r['errors']) and st['failed']==0,{k:st[k] for k in ('saved','dupFile','failed')})
        check("Нет ошибок JavaScript (Android)",not errs,errs[:3])
        await b.close()
    n=sum(1 for _,ok in results if ok);print(f"\nИТОГО: {n} из {len(results)} проверок пройдено")
    json.dump(results,open(f"{OUT}/results.json","w"),ensure_ascii=False)
asyncio.run(main())
