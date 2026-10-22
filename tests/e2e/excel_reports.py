import os, sys as _s;_s.path.insert(0,os.path.dirname(os.path.abspath(__file__)));from authhelp import login
FIX=os.environ.get("VYP_FIX","fixtures-out");OUTDIR=os.environ.get("VYP_OUT","/tmp/vypiski-test-out")
import asyncio, sys, os, glob, json, csv, io, http.server, threading, functools
from playwright.async_api import async_playwright
ROOT=sys.argv[1]; OUT=f'{OUTDIR}/xlout'; os.makedirs(OUT,exist_ok=True)
files=sorted(glob.glob(f'{FIX}/xlcases/*.docx'))+sorted(glob.glob(f'{FIX}/orders/*.docx'))
h=functools.partial(http.server.SimpleHTTPRequestHandler, directory=ROOT); h.log_message=lambda *a:None
srv=http.server.ThreadingHTTPServer(("127.0.0.1",0),h); threading.Thread(target=srv.serve_forever,daemon=True).start()
res=[]
def check(n,ok,d=''):res.append(ok);print(('PASS ' if ok else 'FAIL ')+n+(' — '+str(d) if d else ''))
async def main():
    async with async_playwright() as p:
        b=await p.chromium.launch()
        for width in [1366,390]:
            ctx=await b.new_context(viewport={"width":width,"height":900},locale="ru-RU",accept_downloads=True);pg=await ctx.new_page()
            errs=[];pg.on("pageerror",lambda e:errs.append(str(e)));pg.on("dialog",lambda d:asyncio.ensure_future(d.accept('Мой отчёт')))
            await pg.goto(f"http://127.0.0.1:{srv.server_address[1]}/index.html");await pg.wait_for_timeout(1500)
            await login(pg)
            await pg.evaluate("()=>document.querySelector(`[data-sec='import']`).click()");await pg.wait_for_timeout(500)
            await pg.set_input_files("#wordFiles",files);await pg.wait_for_timeout(1500)
            await pg.evaluate("()=>[...document.querySelectorAll('#content button.primary')].find(x=>/Импорт/.test(x.textContent)).click()")
            for _ in range(100):
                await pg.wait_for_timeout(300)
                if not await pg.evaluate("()=>VypiskiImportJob.state().running"):break
            await pg.evaluate("()=>document.querySelector(`#appNav [data-sec='table']`).click()")
            for _ in range(80):
                await pg.wait_for_timeout(250)
                if 'Разобрано' in await pg.inner_text('#rg .rg-status'):break
            if width==1366:
                check('Пункт меню «Excel» и заголовок «Excel-аналитика»',await pg.inner_text('#appNav [data-sec=table]')=='Excel' and 'Excel-аналитика' in await pg.inner_text('#rg h2'))
                await pg.click("#rg .rg-group .rg-chip:has-text('Отпуск')");await pg.wait_for_timeout(500)
                await pg.click("#rg .rg-chips button:has-text('2024')");await pg.wait_for_timeout(300);await pg.click("#rg .rg-chips button:has-text('2025')");await pg.wait_for_timeout(500)
                title=await pg.input_value('#rg .rg-title');check('Заголовок как в образце',title=='Все приказы по отпускам с 2024 по 2025 г.',title)
                rows=await pg.evaluate("()=>[...document.querySelectorAll('#rg tbody tr:not(.grp)')].map(tr=>[...tr.children].map(td=>td.innerText.trim()))")
                fio=[r[2] for r in rows];check('Отпуска 2024–2025: Иванов, Елисеев, Иванченко',{'ИВАНОВ И.И.','ЕЛИСЕЕВ Р.А.','ИВАНЧЕНКО Р.А.'}<=set(fio),fio)
                el=[r for r in rows if r[2]=='ЕЛИСЕЕВ Р.А.' and '№ 1'==r[4]];check('Елисеев: основной за 2024 и дополнительный за 2023 в одной строке',el and 'основной отпуск за 2024 и дополнительный за 2023' in el[0][8],el[:1])
                canc=[r for r in rows if 'отмена' in r[7]];check('Отмена отпуска отличается от предоставления',canc and 'приказ № 1 от 21.01.2024' in canc[0][7],canc[:1])
                await pg.screenshot(path=f'{OUT}/desk-leaves.png')
                # months
                await pg.click("#rg .rg-chips button[title='март']");await pg.wait_for_timeout(400)
                n=await pg.evaluate("()=>[...document.querySelectorAll('#rg tbody tr:not(.grp)')].map(tr=>tr.children[5].innerText)")
                check('Фильтр по месяцу (март)',n and all(x.split('.')[1]=='03' for x in n),n)
                await pg.click("#rg .rg-chips button[title='март']");await pg.wait_for_timeout(300)
                # sort by FIO
                await pg.click("#rg th:has-text('ФИО')");await pg.wait_for_timeout(300)
                f2=await pg.evaluate("()=>[...document.querySelectorAll('#rg tbody tr:not(.grp)')].map(tr=>tr.children[2].innerText)")
                check('Сортировка по клику на заголовок',f2==sorted(f2,key=lambda x:x),f2[:4])
                # export XLSX
                async with pg.expect_download() as d:
                    await pg.click("#rg button:has-text('Скачать XLSX')")
                dl=await d.value;xp=f'{OUT}/{dl.suggested_filename}';await dl.save_as(xp)
                import openpyxl
                wb=openpyxl.load_workbook(xp);check('XLSX: листы Данные / Параметры / На проверку',wb.sheetnames[:2]==['Данные','Параметры'],wb.sheetnames)
                ws=wb['Данные'];check('XLSX: автофильтр и закреплённая шапка',ws.auto_filter.ref and ws.auto_filter.ref.startswith('A3') and ws.freeze_panes=='A4',(ws.auto_filter.ref,ws.freeze_panes))
                check('XLSX: даты — настоящие даты',type(ws['F4'].value).__name__=='datetime',type(ws['F4'].value))
                check('XLSX: итоги и период',any('Итого' in str(c.value) for c in ws['A']) and 'Период' in str(ws['A2'].value),ws['A2'].value)
                # CSV with formula guard
                await pg.click("#rg button:has-text('Сбросить')");await pg.wait_for_timeout(800)
                await pg.fill("#rg input[placeholder='Например: Иванов или Иванов И.И.']","Зотов");await pg.wait_for_timeout(400)
                async with pg.expect_download() as d:
                    await pg.click("#rg button:has-text('CSV')")
                dl=await d.value;cp=f'{OUT}/{dl.suggested_filename}';await dl.save_as(cp)
                raw=open(cp,'rb').read();txt=raw.decode('utf-8-sig')
                check('CSV: UTF-8 с BOM',raw[:3]==b'\xef\xbb\xbf')
                check('CSV: формула «=СУММ…» обезврежена',"\"'=СУММ" in txt and '"=СУММ' not in txt)
                # person → all events
                await pg.fill("#rg input[placeholder='Например: Иванов или Иванов И.И.']","");await pg.wait_for_timeout(300)
                await pg.click("#rg td:has-text('ИВАНОВ И.И.') >> nth=0");await pg.wait_for_timeout(400)
                await pg.click("#rgSheet button:has-text('Все события этого человека')");await pg.wait_for_timeout(800)
                ev=await pg.evaluate("()=>[...document.querySelectorAll('#rg tbody tr:not(.grp)')].map(tr=>tr.children[6].innerText+'|'+tr.children[7].innerText)")
                check('Все события одного человека (ФИО → события)',len(ev)>=2,ev[:5])
                # manual correction
                await pg.click("#rg tbody tr:not(.grp) >> nth=0");await pg.wait_for_timeout(300)
                await pg.click("#rgSheet summary:has-text('Исправить')");await pg.fill("#rgSheet [data-k='detail']","исправлено вручную: тест")
                await pg.fill('#rgSheet .fix input[type=text]:not([data-k])','сверено с оригиналом');await pg.click("#rgSheet button:has-text('Сохранить исправление')");await pg.wait_for_timeout(1200)
                t=await pg.inner_text('#rg');check('Ручная коррекция сохраняется и помечается',"исправлено вручную: тест" in t and 'исправлено' in t)
                same=await pg.evaluate("async()=>{const o=(await VypiskiRepository.all('orders')).find(o=>o.sourceFile==='Приказ_1_2024.docx');return o.textSha256===await VypiskiImport.sha256(VypiskiImport.normText(o.text))}")
                check('Оригинал приказа не изменился после коррекции',same)
                # group by person + selection
                await pg.click("#rg button:has-text('Сбросить')");await pg.wait_for_timeout(800)
                await pg.evaluate("()=>{const s=[...document.querySelectorAll('#rg select')].find(x=>[...x.options].some(o=>o.value==='fio'));s.value='fio';s.dispatchEvent(new Event('change'))}");await pg.wait_for_timeout(500)
                check('Группировка по ФИО',await pg.locator('#rg tr.grp').count()>3)
                await pg.click('#rg tbody tr:not(.grp) >> nth=0 >> input.rg-ck');await pg.click('#rg tbody tr:not(.grp) >> nth=1 >> input.rg-ck')
                check('Выбор строк и панель «Создать выписки»',await pg.is_visible('#rg .rg-bulk') and 'Выбрано строк: 2' in await pg.inner_text('#rg .rg-bulk'))
                await pg.evaluate("()=>{const s=[...document.querySelectorAll('#rg select')].find(x=>[...x.options].some(o=>o.value==='fio'));s.value='';s.dispatchEvent(new Event('change'))}")
                await pg.screenshot(path=f'{OUT}/desk-all.png')
            else:
                await pg.click("#rg .rg-group .rg-chip:has-text('Отпуск')");await pg.wait_for_timeout(600)
                check('Телефон: вид «Отпуск», карточки',await pg.locator('#rg .rg-card').count()>0 and await pg.locator('#rg .rg-tablewrap').is_hidden())
                await pg.screenshot(path=f'{OUT}/phone-check.png',full_page=False)
            check(f'Нет ошибок JavaScript ({width}px)',not errs,errs[:3])
            await ctx.close()
        await b.close()
    print(f'ИТОГО: {sum(res)} из {len(res)}')
asyncio.run(main())
