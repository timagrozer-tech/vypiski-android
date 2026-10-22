import os, sys as _s;_s.path.insert(0,os.path.dirname(os.path.abspath(__file__)));from authhelp import login
FIX=os.environ.get("VYP_FIX","fixtures-out");OUTDIR=os.environ.get("VYP_OUT","/tmp/vypiski-test-out")
import asyncio, sys, os, glob, http.server, threading, functools
from playwright.async_api import async_playwright
import openpyxl
ROOT=sys.argv[1]; OUT=f'{OUTDIR}/xpout'; os.makedirs(OUT,exist_ok=True)
files=sorted(glob.glob(f'{FIX}/xlcases/*.docx'))+sorted(glob.glob(f'{FIX}/orders/*.docx'))
XP=f'{FIX}/xpcases/'
h=functools.partial(http.server.SimpleHTTPRequestHandler, directory=ROOT); h.log_message=lambda *a:None
srv=http.server.ThreadingHTTPServer(("127.0.0.1",0),h); threading.Thread(target=srv.serve_forever,daemon=True).start()
res=[]
def check(n,ok,d=''):res.append(bool(ok));print(('PASS ' if ok else 'FAIL ')+n+(' — '+str(d) if d else ''))
async def main():
  async with async_playwright() as p:
    b=await p.chromium.launch()
    for width in [1366,390]:
      ctx=await b.new_context(viewport={"width":width,"height":900},locale="ru-RU",accept_downloads=True);pg=await ctx.new_page()
      errs=[];pg.on("pageerror",lambda e:errs.append(str(e)))
      await pg.goto(f"http://127.0.0.1:{srv.server_address[1]}/index.html");await login(pg)
      await pg.evaluate("()=>document.querySelector(`[data-sec='import']`).click()");await pg.wait_for_timeout(400)
      await pg.set_input_files("#wordFiles",files);await pg.wait_for_timeout(1500)
      await pg.evaluate("()=>[...document.querySelectorAll('#content button.primary')].find(x=>/Импорт/.test(x.textContent)).click()")
      for _ in range(100):
        await pg.wait_for_timeout(300)
        if not await pg.evaluate("()=>VypiskiImportJob.state().running"):break
      before=await pg.evaluate("async()=>{const o=await VypiskiRepository.all('orders');return Object.fromEntries(o.map(x=>[x.id,x.textSha256+'|'+x.fileSha256]))}")
      await pg.evaluate("()=>VypiskiRegisterView.open()");await pg.wait_for_selector('#rg .xp-tabs')
      tabs=await pg.locator('#rg .xp-tabs button').all_inner_texts();check(f'[{width}] Вкладки Excel',tabs==['Отчёты','Импорт и сверка','Шаблоны'],tabs)
      check(f'[{width}] Нет кнопок, повторяющих вкладки',await pg.locator("#rg .xp-quick").count()==0)
      await pg.wait_for_function("()=>/Индекс обновлён: \\d/.test(document.querySelector('#rg .xp-stat')?.textContent||'')",timeout=20000)
      check(f'[{width}] Панель статуса',True,await pg.inner_text('#rg .xp-stat'))
      await pg.click("#rg .xp-tabs button:has-text('Импорт и сверка')")
      vis=await pg.evaluate("()=>getComputedStyle(document.querySelector('#rg .rg-actions')).display");check(f'[{width}] Отчёты скрыты на вкладке импорта',vis=='none')
      # XLS refused
      await pg.set_input_files("#rg .xp-pane input[type=file]",XP+'old.xls');await pg.wait_for_selector('#rg .xp-bad')
      t=await pg.inner_text('#rg .xp-bad');check(f'[{width}] XLS: понятный отказ и путь конвертации','XLS' in t and 'xlsx' in t.lower() and 'облач' in t.lower(),t)
      # XLSM with macro
      await pg.set_input_files("#rg .xp-pane input[type=file]",XP+'svod_macro.xlsm');await pg.wait_for_selector('#rg .xp-pane .xp-prev table')
      w=' '.join(await pg.locator('#rg .xp-pane .xp-warn').all_inner_texts())
      check(f'[{width}] Макросы, внешние ссылки, формулы — предупреждения, не исполняются','макрос' in w.lower() and 'внешние ссылки' in w and 'формул' in w,w[:200])
      hr=await pg.input_value("#rg .xp-pane input[type=number]");check(f'[{width}] Строка заголовков найдена (после объединённого заголовка)',hr=='2',hr)
      sel=await pg.evaluate("()=>[...document.querySelectorAll('#rg .xp-pane .xp-grid label')].filter(l=>l.querySelector('select')).map(l=>l.firstChild.textContent+'='+l.querySelector('select').value)")
      check(f'[{width}] Сопоставление столбцов автоматически',all(x in sel for x in ['ФИО=1','Звание=2','№ приказа=3','Дата приказа=4','Вид события=5','Начало=6','Окончание=7']),sel)
      await pg.click("#rg .xp-pane button:has-text('Сверить с приказами')");await pg.wait_for_selector('#rg .xp-sum button',timeout=30000);await pg.wait_for_timeout(300)
      rows=await pg.evaluate("()=>[...document.querySelectorAll('#rg .xp-res .xp-row')].map(r=>({line:r.querySelector('.xl .h').textContent,st:r.querySelector('.xs').textContent,t:r.innerText}))")
      st={r['line'].split()[-1]:r['st'] for r in rows}
      exp={'3':'расхождение','4':'совпадает','5':'неоднозначно','6':'нет подтверждения','7':'расхождение','8':'совпадает'}
      check(f'[{width}] Статусы сверки по строкам',all(st.get(k)==v for k,v in exp.items()),st)
      t3=[r['t'] for r in rows if r['line'].endswith(' 3')][0];check(f'[{width}] Расхождение: оба значения рядом','04.02.2025' in t3 and '03.02.2025' in t3,t3[-200:])
      t7=[r['t'] for r in rows if r['line'].endswith(' 7')][0];check(f'[{width}] Ведущие нули номера сохранены и показаны','0178' in t7 and 'ведущих нулей' in t7)
      t8=[r['t'] for r in rows if r['line'].endswith(' 8')][0];check(f'[{width}] Некорректная дата 31.02.2025 отмечена','некорректная дата' in t8)
      t5=[r['t'] for r in rows if r['line'].endswith(' 5')][0];check(f'[{width}] Однофамильцы не объединяются','Иванов Иван Иванович' in t5 and 'Иванов Игорь Ильич' in t5)
      check(f'[{width}] Ссылка на первоисточник',await pg.locator("#rg .xp-res button:has-text('Открыть первоисточник')").count()>=3)
      await pg.screenshot(path=f'{OUT}/rec-{width}.png',full_page=False)
      async with pg.expect_download() as d: await pg.click("#rg .xp-pane button:has-text('Скачать результат сверки')")
      dl=await d.value;rp=f'{OUT}/{width}-rec.xlsx';await dl.save_as(rp);wb=openpyxl.load_workbook(rp);ws=wb['Сверка']
      check(f'[{width}] Результат сверки: новый XLSX с фильтром и шапкой',ws.auto_filter.ref and ws.freeze_panes=='A4' and ws.max_row==3+len(rows),(ws.auto_filter.ref,ws.freeze_panes,ws.max_row))
      check(f'[{width}] Дата в результате — настоящая дата',any(type(c.value).__name__=='datetime' for c in ws['E'][3:]))
      # CSV 1251 ;
      await pg.set_input_files("#rg .xp-pane input[type=file]",XP+'svod_1251.csv');await pg.wait_for_selector('#rg .xp-pane .xp-prev table');await pg.wait_for_timeout(200)
      enc=await pg.evaluate("()=>[...document.querySelectorAll('#rg .xp-pane .xp-grid select')].slice(0,2).map(s=>s.value)")
      check(f'[{width}] CSV: Windows-1251 и «;» определены',enc==['windows-1251',';'],enc)
      pv=await pg.inner_text('#rg .xp-pane .xp-prev');check(f'[{width}] CSV: кириллица читается','Петров Сергей' in pv)
      await pg.click("#rg .xp-pane button:has-text('Сверить с приказами')");await pg.wait_for_timeout(1500)
      async with pg.expect_download() as d: await pg.click("#rg .xp-pane button:has-text('Скачать результат сверки')")
      dl=await d.value;rp=f'{OUT}/{width}-rec-csv.xlsx';await dl.save_as(rp);ws=openpyxl.load_workbook(rp)['Сверка']
      f=[c for c in ws['B'] if c.value and 'HYPERLINK' in str(c.value)]
      check(f'[{width}] =HYPERLINK из таблицы остаётся текстом',f and f[0].data_type=='s',[(c.value,c.data_type) for c in f])
      await pg.set_input_files("#rg .xp-pane input[type=file]",XP+'svod_utf8_comma.csv');await pg.wait_for_selector('#rg .xp-pane .xp-prev table');await pg.wait_for_timeout(200)
      enc=await pg.evaluate("()=>[...document.querySelectorAll('#rg .xp-pane .xp-grid select')].slice(0,2).map(s=>s.value)")
      check(f'[{width}] CSV: UTF-8 и «,» определены',enc==['utf-8',','],enc)
      # reports: split over the row limit and by kind
      await pg.click("#rg .xp-tabs button:has-text('Отчёты')");await pg.evaluate("()=>VypiskiRegisterView._rowLimit(5)")
      await pg.click("#rg .rg-chips button:has-text('Весь период')");await pg.wait_for_selector('#rg .xp-tabs');await pg.click("#rg .xp-tabs button:has-text('Отчёты')")
      async with pg.expect_download() as d: await pg.click("#rg button:has-text('Скачать XLSX')")
      dl=await d.value;rp=f'{OUT}/{width}-split.xlsx';await dl.save_as(rp);wb=openpyxl.load_workbook(rp)
      note=await pg.inner_text('#rg .rg-xlnote')
      check(f'[{width}] Лимит строк: деление на листы с уведомлением',sum(1 for n in wb.sheetnames if n.startswith('Данные'))>=2 and 'разделены' in note,(wb.sheetnames[:5],note[:120]))
      ok=all(wb[n].max_row<=5+3+2 for n in wb.sheetnames if n.startswith('Данные'));check(f'[{width}] Ни один лист не превышает лимит',ok)
      await pg.evaluate("()=>VypiskiRegisterView._rowLimit(0)")
      await pg.evaluate("()=>{localStorage.setItem('vypiski.xl.split','category')}")
      async with pg.expect_download() as d: await pg.click("#rg button:has-text('Скачать XLSX')")
      dl=await d.value;rp=f'{OUT}/{width}-bycat.xlsx';await dl.save_as(rp);wb=openpyxl.load_workbook(rp)
      check(f'[{width}] Лист на каждый вид события','Отпуск' in wb.sheetnames and 'Поощрение' in wb.sheetnames,wb.sheetnames)
      await pg.evaluate("()=>{localStorage.setItem('vypiski.xl.split','')}")
      # templates tab
      await pg.click("#rg .xp-tabs button:has-text('Шаблоны')");await pg.wait_for_selector("#rg .xp-pane[data-pane=tpl] .xp-list")
      await pg.wait_for_timeout(500);hist=await pg.inner_text("#rg .xp-pane[data-pane=tpl]")
      check(f'[{width}] Шаблоны и история выгрузок с SHA-256','Отпуска за год' in hist and 'SHA-256' in hist and 'admin' in hist,hist[-300:])
      await pg.click("#rg .xp-pane[data-pane=tpl] .it:has-text('Поощрения') button:has-text('Сформировать')");await pg.wait_for_timeout(1200)
      t=await pg.input_value('#rg .rg-title');check(f'[{width}] Запуск шаблона пересчитывает отчёт','поощрен' in t.lower(),t)
      after=await pg.evaluate("async()=>{const o=await VypiskiRepository.all('orders');return Object.fromEntries(o.map(x=>[x.id,x.textSha256+'|'+x.fileSha256]))}")
      check(f'[{width}] SHA-256 приказов не изменились',before==after)
      check(f'[{width}] Без ошибок JS',not errs,errs[:3])
      await ctx.close()
    await b.close()
  print(f'{sum(res)}/{len(res)}')
asyncio.run(main())
