import sys,os
from docx import Document
OUT=(sys.argv[1] if len(sys.argv)>1 else 'smpcases').rstrip('/')+'/';os.makedirs(OUT,exist_ok=True)
from docx.shared import Pt,Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
def mk(path,head,body,certify=True,size=12):
  d=Document();s=d.sections[0];s.top_margin=Mm(20);s.bottom_margin=Mm(20);s.left_margin=Mm(30);s.right_margin=Mm(15)
  def P(t,al=A.JUSTIFY,b=False,fl=None):
    p=d.add_paragraph();r=p.add_run(t);r.font.size=Pt(size);r.bold=b;p.alignment=al
    if fl is not None:p.paragraph_format.first_line_indent=Mm(fl)
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
  P('ВЫПИСКА',A.CENTER,True)
  for h in head:P(h,A.CENTER)
  for t in body:P(t,A.JUSTIFY,False,12.5)
  if certify:
    for t in ['ВЕРНО:','Начальник строевой части','капитан','И.Петров']:P(t,A.LEFT)
  d.save(path)
H=['из приказа командира войсковой части 00000','«06» июня 2025 г. № 158','г. Москва']
B=['§ 5','Нижепоименованным военнослужащим произвести выплату:','31. Младшему сержанту Елисееву Дмитрию Алексеевичу, СУ-067677, командиру стрелкового отделения.','Основание: рапорт начальника физической подготовки от 17 апреля 2025 года № 11090.']
mk(OUT+'Выписка_158_Елисеев.docx',H,B)
B2=list(B);B2[2]='31. Младшему сержанту Елисееву Дмитрию Алексеевичу, СУ-067677, командиру разведывательного отделения.'
mk(OUT+'Выписка_158_с_ошибкой.docx',['из приказа командира войсковой части 00000','от 06.06.2025 № 158'],B2)
mk(OUT+'Выписка_999_нет.docx',['из приказа командира войсковой части 00000','«01» января 2024 г. № 999'],['§ 1','1. Рядовому Неизвестному Ивану Ивановичу, АА-000000, стрелку.'])
# no number in header: found by personal number
mk(OUT+'Выписка_без_номера.docx',['из приказа командира войсковой части 00000'],B[1:3],certify=False)
