"""A long synthetic order (invented names): 5 § with 60 points each, two lines per point — for the speed check of bulk preparation."""
import sys,os
from docx import Document
OUT=(sys.argv[1] if len(sys.argv)>1 else 'bigcase').rstrip('/')+'/';os.makedirs(OUT,exist_ok=True)
SN=['Арбузов','Брусков','Вишняков','Грушин','Дынин','Ежевикин','Жасминов','Земляникин','Ирисов','Калинин','Лимонов','Малинин','Нарциссов','Орехов','Персиков','Рябинин','Сливин','Тыквин','Укропов','Фисташкин']
NM=[('Иван','Ивана'),('Пётр','Петра'),('Олег','Олега'),('Денис','Дениса'),('Роман','Романа')]
PT=[('Ильич','Ильича'),('Петрович','Петровича'),('Олегович','Олеговича')]
d=Document()
for t in ['ПРИКАЗ','КОМАНДИРА ВОЙСКОВОЙ ЧАСТИ 77777','по строевой части','« 12 » ноября 2026 г. № 500','г. Энск','']:d.add_paragraph(t)
k=0
for s in range(1,6):
  d.add_paragraph(f'§ {s}');d.add_paragraph('Полагать убывшими в командировку:')
  for i in range(1,61):
    sn=SN[k%20];nm=NM[(k//20)%5];pt=PT[(k//100)%3];k+=1
    d.add_paragraph(f'{i}. Рядового {sn}а {nm[1]} {pt[1]}, ББ-{300000+k}, стрелка, в г. Тверь с 13 ноября 2026 года.')
    d.add_paragraph(f'Основание: рапорт № {k}.')
for t in ['','КОМАНДИР ВОЙСКОВОЙ ЧАСТИ 77777','полковник\tА.Энский']:d.add_paragraph(t)
d.save(OUT+'Приказ 500 от 12.11.2026.docx')
