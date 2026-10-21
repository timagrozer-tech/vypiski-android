"""Synthetic test data for Выписки (invented names, no personal data). Usage: python3 make_fixtures.py <out-dir>"""
import os, sys, io, random, zipfile, shutil, docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
OUT=os.path.abspath(sys.argv[1] if len(sys.argv)>1 else 'fixtures-out'); os.makedirs(OUT,exist_ok=True)

def orders(out):
    os.makedirs(out,exist_ok=True)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    people = [("Иванову Ивану Ивановичу", "Иванов Иван Иванович"), ("Петрову Сергею Николаевичу", "Петров Сергей Николаевич"),
              ("Сидоровой Анне Павловне", "Сидорова Анна Павловна"), ("Кузнецову Андрею Олеговичу", "Кузнецов Андрей Олегович")]
    random.seed(1)
    for n in range(1, 9):
        d = docx.Document()
        st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(14)
        for line in ["КОМАНДИР ВОЙСКОВОЙ ЧАСТИ 00000", "ПРИКАЗ", "(по строевой части)", f"«{n+3}» марта 2016 г. № {100+n}", "г. Москва"]:
            p = d.add_paragraph(line); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for s in range(1, 4):
            d.add_paragraph(f"§ {s}").alignment = WD_ALIGN_PARAGRAPH.CENTER
            dat, nom = random.choice(people)
            verb = random.choice(["Назначить на должность заместителя командира роты", "Предоставить основной отпуск за 2016 год продолжительностью 30 суток", "Исключить из списков личного состава части", "Объявить благодарность за образцовое исполнение обязанностей"])
            d.add_paragraph(f"{s}. {verb} рядового {dat}, личный номер А-{random.randint(100000,999999)}.")
            d.add_paragraph(f"Основание: рапорт {nom.split()[0]}а от {n}.03.2016 г.")
        d.add_paragraph("")
        d.add_paragraph("КОМАНДИР ВОЙСКОВОЙ ЧАСТИ 00000")
        d.add_paragraph("полковник                                   А.Смирнов")
        d.save(f"{out}/Приказ_{100+n}_2016.docx")

def many(n,out):
    os.makedirs(out,exist_ok=True)
    random.seed(7)
    names=[("Иванову Ивану Ивановичу","Иванова"),("Петрову Сергею Николаевичу","Петрова"),("Сидоровой Анне Павловне","Сидоровой"),("Кузнецову Андрею Олеговичу","Кузнецова"),("Смирнову Павлу Петровичу","Смирнова"),("Волкову Олегу Игоревичу","Волкова")]
    verbs=["Назначить на должность заместителя командира роты","Предоставить основной отпуск за 2016 год продолжительностью 30 суток","Исключить из списков личного состава части","Объявить благодарность за образцовое исполнение обязанностей","Зачислить в списки личного состава части и на все виды обеспечения","Полагать сдавшим дела и должность"]
    base=docx.Document()
    for k in range(n):
        d=docx.Document(); st=d.styles["Normal"]; st.font.name="Times New Roman"; st.font.size=Pt(14)
        m=k%12+1; day=k%27+1
        for line in ["КОМАНДИР ВОЙСКОВОЙ ЧАСТИ 00000","ПРИКАЗ","(по строевой части)",f"«{day}» {['января','февраля','марта','апреля','мая','июня','июля','августа','сентября','октября','ноября','декабря'][m-1]} 2016 г. № {k+1}","г. Москва"]:
            d.add_paragraph(line)
        for s in range(1, random.randint(4,12)):
            d.add_paragraph(f"§ {s}")
            for q in range(random.randint(1,3)):
                dat,gen=random.choice(names)
                d.add_paragraph(f"{q+1}. {random.choice(verbs)} рядового {dat}, личный номер А-{random.randint(100000,999999)}. "*random.randint(1,3))
                d.add_paragraph(f"Основание: рапорт {gen} от {day}.{m:02d}.2016 г.")
        if k%5==0:
            t=d.add_table(rows=3,cols=3)
            for r in range(3):
                for c in range(3): t.cell(r,c).text=f"Ячейка {r}{c}"
        d.add_paragraph("КОМАНДИР ВОЙСКОВОЙ ЧАСТИ 00000"); d.add_paragraph("полковник\tА.Смирнов")
        sub=f"{out}/2016/{m:02d}"; os.makedirs(sub,exist_ok=True)
        d.save(f"{sub}/Приказ_{k+1}.docx")

def edge(out):
    os.makedirs(out,exist_ok=True)
    d=docx.Document(); d.add_paragraph("seed"); buf=io.BytesIO(); d.save(buf); base=buf.getvalue()
    W='xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    EXTRA='xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    def mk(name, body, decl=True, root=None):
        zin=zipfile.ZipFile(io.BytesIO(base)); zout=io.BytesIO(); z=zipfile.ZipFile(zout,"w",zipfile.ZIP_DEFLATED)
        for it in zin.infolist():
            data=zin.read(it.filename)
            if it.filename=="word/document.xml":
                x=(('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n') if decl else '')+(root or f'<w:document {W} {EXTRA}><w:body>{body}<w:sectPr/></w:body></w:document>')
                data=x.encode("utf-8")
            z.writestr(it,data)
        z.close(); open(f"{out}/{name}.docx","wb").write(zout.getvalue())
    P=lambda t:f'<w:p><w:r><w:t>{t}</w:t></w:r></w:p>'
    mk("01-basic", P("КОМАНДИР ВОЙСКОВОЙ ЧАСТИ")+P("«4» марта 2016 г. № 101")+P("§ 1"))
    mk("02-tabs-breaks", '<w:p><w:r><w:t xml:space="preserve">полковник </w:t><w:tab/><w:t>А.Смирнов</w:t><w:br/><w:t>вторая</w:t><w:cr/><w:t>третья</w:t><w:noBreakHyphen/><w:t>x</w:t></w:r></w:p>'+P("конец     "))
    mk("03-entities", P("&lt;тег&gt; &amp; &quot;кавычки&quot; &apos;апостроф&apos; &#1046; &#x416; ☺")+'<w:p><w:r><w:t><![CDATA[сырой <текст> & данные]]></w:t></w:r></w:p>')
    mk("04-tracked", '<w:p><w:r><w:t>было </w:t></w:r><w:del w:id="1" w:author="a"><w:r><w:delText>удалено</w:delText></w:r></w:del><w:ins w:id="2" w:author="a"><w:r><w:t>вставлено</w:t></w:r></w:ins></w:p>'
       '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText xml:space="preserve"> PAGE </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>7</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>')
    tbl=lambda rows:'<w:tbl><w:tblPr/>'+''.join('<w:tr>'+''.join(rows_c for rows_c in r)+'</w:tr>' for r in rows)+'</w:tbl>'
    tc=lambda t,vm=None:'<w:tc><w:tcPr>'+({None:'','restart':'<w:vMerge w:val="restart"/>','cont':'<w:vMerge/>','contv':'<w:vMerge w:val="continue"/>'}[vm])+'</w:tcPr>'+''.join(P(x) for x in t.split("|"))+'</w:tc>'
    mk("05-table-merge", tbl([[tc("A",'restart'),tc("B")],[tc("A2",'cont'),tc("B2|B3")],[tc("A3",'contv'),tc("")]]))
    inner=tbl([[tc("in1"),tc("in2")]])
    mk("06-nested-table", '<w:tbl><w:tr><w:tc>'+P("outer")+inner+P("after")+'</w:tc><w:tc>'+P("x")+'</w:tc></w:tr></w:tbl>'+P("tail"))
    mk("07-sdt", '<w:sdt><w:sdtPr/><w:sdtContent>'+P("в контроле 1")+P("в контроле 2")+tbl([[tc("т в sdt")]])+'</w:sdtContent></w:sdt>'+P("после"))
    mk("08-textbox", '<w:p><w:r><w:t>до рамки </w:t></w:r><w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData><wps:wsp><wps:txbx><w:txbxContent>'+P("в рамке")+'</w:txbxContent></wps:txbx></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice><mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>'+P("в рамке")+'</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r><w:r><w:t> после рамки</w:t></w:r></w:p>')
    mk("09-numbering-hyperlink", '<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr><w:r><w:t>пункт списка</w:t></w:r></w:p><w:p><w:hyperlink r:id="x"><w:r><w:t>ссылка</w:t></w:r></w:hyperlink><w:bookmarkStart w:id="0" w:name="b"/><w:r><w:t> закладка</w:t></w:r><w:bookmarkEnd w:id="0"/></w:p>')
    mk("10-other-prefix", None, root='<ns0:document xmlns:ns0="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><ns0:body><ns0:p><ns0:r><ns0:t>чужой префикс</ns0:t></ns0:r></ns0:p><ns0:tbl><ns0:tr><ns0:tc><ns0:tcPr><ns0:vMerge ns0:val="continue"/></ns0:tcPr><ns0:p><ns0:r><ns0:t>c</ns0:t></ns0:r></ns0:p></ns0:tc></ns0:tr></ns0:tbl></ns0:body></ns0:document>')
    mk("11-default-ns", None, root='<document xmlns="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><body><p><r><t>по умолчанию</t></r></p></body></document>')
    mk("12-empty", "", )
    mk("13-broken", None, root=f'<w:document {W}><w:body><w:p><w:r><w:t>сломан</w:r></w:p></w:body></w:document>')
    mk("14-comment-pi", '<!-- комментарий --><?mso-x y?>'+P("текст")+'<w:p/>'+'<w:p><w:r><w:t xml:space="preserve">   </w:t></w:r></w:p>'+P("ещё"))
    mk("15-attr-gt", '<w:p><w:r><w:t>a</w:t></w:r></w:p><w:customXml w:element="x&gt;y" w:uri="a>b"><w:p><w:r><w:t>внутри customXml</w:t></w:r></w:p></w:customXml>'+P("b"))
    mk("16-nodecl-ws", "\n  "+P("с переводами строк")+"\n  ", decl=False)
    mk("17-bad-entity", P("плохая &nbsp; сущность"))

def accept_zip(src_dir,path):
    src=sorted(f for f in os.listdir(src_dir) if f.endswith('.docx')); src=[os.path.join(src_dir,f) for f in src]
    z=zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED)
    for folder,fs in [('приказы/2025/Январь/',src[0:3]),('приказы/2025/Февраль/',src[3:5]),('приказы/2026/Март/',src[5:8])]:
        for f in fs: z.write(f,folder+os.path.basename(f))
    z.write(src[0],'приказы/2026/Март/копия/'+os.path.basename(src[0]))
    z.writestr('приказы/2025/Январь/скан.pdf',b'%PDF-1.4 fake')
    z.writestr('приказы/2025/Январь/~$Приказ_101_2016.docx',b'x'*162)
    inner=io.BytesIO();zz=zipfile.ZipFile(inner,'w');zz.writestr('a.docx',b'x');zz.close();z.writestr('приказы/вложенный.zip',inner.getvalue())
    z.writestr('приказы/../evil.docx',open(src[1],'rb').read())
    z.writestr(zipfile.ZipInfo('приказы/2026/бомба.docx'),b'\0'*(60*1024*1024+5),compress_type=zipfile.ZIP_DEFLATED)
    z.close()

orders(f"{OUT}/orders"); many(int(os.environ.get('MANY','300')),f"{OUT}/many"); edge(f"{OUT}/edge"); os.makedirs(f"{OUT}/accept",exist_ok=True); accept_zip(f"{OUT}/orders",f"{OUT}/accept/приказы.zip")
print('fixtures in',OUT)
